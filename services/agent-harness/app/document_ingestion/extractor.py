from __future__ import annotations

import io
import json
import re
import zipfile
from dataclasses import asdict
from typing import Any

from app.document_ingestion.models import Candidate, ExtractedDocument, ExtractionResult, IngestionJob


QUESTION_RE = re.compile(r"^(?:[-*]\s*)?(?:Q\d+[:.)]\s*)?(?P<text>[^?\n]{6,220}\?)\s*$", re.IGNORECASE)
CUE_RE = re.compile(r"^(?:describe|talk about|speak about)\s+.+", re.IGNORECASE)
FOLLOWUP_RE = re.compile(r"^(?:follow[- ]?up|part\s*3|追问|延伸问题)[:：]?\s*(?P<text>.+\?)", re.IGNORECASE)
HEADING_RE = re.compile(r"^\s{0,3}#{1,6}\s+(?P<title>.+?)\s*$")
BACKGROUND_FIELD_RE = re.compile(
    r"^(?P<key>name|目标分数|target band|occupation|job|职业|major|专业|school|学校|hobby|兴趣|weakness|弱项|strength|强项|city|城市)\s*[:：]\s*(?P<value>.+)$",
    re.IGNORECASE,
)


class DocumentParseError(RuntimeError):
    pass


def strip_nul_bytes(text: str) -> str:
    """Strip NUL bytes (0x00) that PostgreSQL text fields cannot contain."""
    return text.replace("\x00", "")


def extract_document(content: bytes, *, filename: str, extension: str, mime_type: str) -> ExtractedDocument:
    extension = extension.lower().strip()
    if extension in {".md", ".markdown", ".txt"}:
        return extract_plain_text(content, filename=filename, extension=extension)
    if extension == ".pdf" or mime_type == "application/pdf":
        return extract_pdf_text(content, filename=filename)
    if extension == ".docx" or mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_docx_text(content, filename=filename)
    raise DocumentParseError(f"unsupported document type: {extension or mime_type}")


def extract_plain_text(content: bytes, *, filename: str, extension: str) -> ExtractedDocument:
    text = strip_nul_bytes(decode_text(content))
    parser_name = "markdown_plain_text" if extension in {".md", ".markdown"} else "plain_text"
    return ExtractedDocument(
        text=text,
        parser_name=parser_name,
        command_log=f"{parser_name}: decoded {len(content)} bytes from {filename}",
        metadata={"filename": filename, "byte_size": len(content)},
    )


def extract_pdf_text(content: bytes, *, filename: str) -> ExtractedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - 依赖缺失只在部署配置错误时触发
        raise DocumentParseError("pypdf dependency is required for PDF ingestion") from exc

    reader = PdfReader(io.BytesIO(content))
    pages: list[str] = []
    for index, page in enumerate(reader.pages):
        page_text = strip_nul_bytes(page.extract_text() or "")
        if page_text.strip():
            pages.append(f"[page {index + 1}]\n{page_text.strip()}")
    return ExtractedDocument(
        text="\n\n".join(pages),
        parser_name="pypdf",
        command_log=f"pypdf: extracted {len(pages)} text pages from {filename}",
        page_count=len(reader.pages),
        metadata={"filename": filename, "page_count": len(reader.pages), "byte_size": len(content)},
    )


def extract_docx_text(content: bytes, *, filename: str) -> ExtractedDocument:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - 依赖缺失只在部署配置错误时触发
        raise DocumentParseError("python-docx dependency is required for DOCX ingestion") from exc

    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise DocumentParseError("invalid DOCX file")
    doc = Document(io.BytesIO(content))
    blocks: list[str] = []
    for paragraph in doc.paragraphs:
        text = strip_nul_bytes(paragraph.text).strip()
        if text:
            blocks.append(text)
    for table in doc.tables:
        for row in table.rows:
            values = [strip_nul_bytes(cell.text).strip() for cell in row.cells if cell.text.strip()]
            if values:
                blocks.append(" | ".join(values))
    return ExtractedDocument(
        text="\n".join(blocks),
        parser_name="python-docx",
        command_log=f"python-docx: extracted {len(blocks)} text blocks from {filename}",
        metadata={"filename": filename, "paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)},
    )


class DocumentSkillRouter:
    def extract(self, job: IngestionJob, document: ExtractedDocument) -> ExtractionResult:
        normalized = normalize_text(document.text)
        skills = select_skills(job, normalized)
        candidates: list[Candidate] = []

        if "question_extractor" in skills:
            candidates.extend(extract_question_candidates(normalized, job))
        if "background_fact_extractor" in skills:
            candidates.extend(extract_background_candidates(normalized, job))
        if "knowledge_chunker" in skills:
            candidates.extend(extract_knowledge_candidates(normalized, job, avoid_duplicate_question_texts=candidates))

        if not candidates and normalized:
            candidates.append(
                Candidate(
                    candidate_kind="knowledge",
                    title=job.source_file.title,
                    summary=first_sentence(normalized),
                    content=normalized[:8000],
                    normalized_payload={"title": job.source_file.title, "content": normalized[:8000]},
                    materialization_plan={"tables": ["knowledge_docs", "knowledge_chunks"], "doc_type": "topic_knowledge"},
                    metadata={"skill": "fallback_document_summary"},
                )
            )

        label, confidence = classify_document(job, normalized, candidates)
        return ExtractionResult(
            classifier_label=label,
            classifier_confidence=confidence,
            candidates=dedupe_candidates(candidates),
            skills_used=skills,
            normalized_text=normalized,
            metadata={
                "parser": document.parser_name,
                "document_metadata": document.metadata,
                "candidate_count": len(candidates),
            },
        )


def select_skills(job: IngestionJob, text: str) -> list[str]:
    requested = job.requested_action
    skills = ["file_sniff", "text_normalizer"]
    if requested in {"auto", "question_bank", "mixed"} or "part 1" in text.lower() or "part 2" in text.lower():
        skills.append("question_extractor")
    if requested in {"auto", "background", "mixed"} or looks_like_background(text):
        skills.append("background_fact_extractor")
    if requested in {"auto", "knowledge", "mixed"} or len(text) > 400:
        skills.append("knowledge_chunker")
    skills.append("materialization_planner")
    return skills


def classify_document(job: IngestionJob, text: str, candidates: list[Candidate]) -> tuple[str, float]:
    counts = {
        "question_bank": sum(1 for item in candidates if item.candidate_kind == "question"),
        "user_profile": sum(1 for item in candidates if item.candidate_kind == "background"),
        "topic_knowledge": sum(1 for item in candidates if item.candidate_kind == "knowledge"),
    }
    if job.requested_action == "question_bank":
        return "question_bank", 0.92
    if job.requested_action == "background":
        return "user_profile", 0.9
    if job.requested_action == "knowledge":
        return "topic_knowledge", 0.88
    winner = max(counts.items(), key=lambda item: item[1])
    if winner[1] == 0:
        return "topic_knowledge", 0.55 if text else 0.2
    total = max(1, sum(counts.values()))
    return winner[0], min(0.95, 0.55 + winner[1] / total * 0.4)


def extract_question_candidates(text: str, job: IngestionJob) -> list[Candidate]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    candidates: list[Candidate] = []
    current_part = part_from_text(text)
    cue_candidate: dict[str, Any] | None = None

    for line in lines:
        lower = line.lower()
        part_match = re.search(r"part\s*([123])", lower)
        if part_match:
            current_part = int(part_match.group(1))

        question_match = QUESTION_RE.match(line)
        followup_match = FOLLOWUP_RE.match(line)
        if followup_match and cue_candidate is not None:
            cue_candidate.setdefault("followups", []).append({"text": followup_match.group("text").strip(), "part": 3})
            continue

        if question_match:
            question_text = question_match.group("text").strip()
            candidates.append(question_candidate(question_text, current_part, job, source_line=line))
            continue

        if CUE_RE.match(line):
            cue_candidate = {
                "text": line.rstrip("."),
                "part": 2,
                "bullets": [],
                "followups": [],
                "source_line": line,
            }
            candidates.append(cue_card_candidate(cue_candidate, job))
            continue

        if cue_candidate is not None and line.startswith(("-", "*")):
            cue_candidate["bullets"].append(line.lstrip("-* ").strip())
            candidates[-1] = cue_card_candidate(cue_candidate, job)

    return candidates[:40]


def question_candidate(question_text: str, part: int, job: IngestionJob, *, source_line: str) -> Candidate:
    payload = {
        "part": part,
        "text": question_text,
        "difficulty": 3,
        "source_type": "user_recall" if job.requested_visibility == "public" else "internal",
    }
    return Candidate(
        candidate_kind="question",
        title=short_title(question_text),
        summary=f"IELTS Speaking Part {part} question",
        content=question_text,
        normalized_payload=payload,
        materialization_plan={"tables": ["questions", "knowledge_docs", "knowledge_chunks"], "part": part},
        metadata={"skill": "question_extractor", "source_line": source_line},
    )


def cue_card_candidate(payload: dict[str, Any], job: IngestionJob) -> Candidate:
    text = str(payload["text"])
    normalized = {
        "part": 2,
        "text": text,
        "cue_card": {
            "prompt": text,
            "bullet_points": payload.get("bullets", []),
            "preparation_seconds": 60,
            "speaking_seconds": 120,
        },
        "followups": payload.get("followups", []),
        "difficulty": 3,
        "source_type": "user_recall" if job.requested_visibility == "public" else "internal",
    }
    content = "\n".join([text, *[f"- {item}" for item in normalized["cue_card"]["bullet_points"]]])
    return Candidate(
        candidate_kind="question",
        title=short_title(text),
        summary="IELTS Speaking Part 2 cue card",
        content=content,
        normalized_payload=normalized,
        materialization_plan={"tables": ["questions", "cue_cards", "followup_templates", "knowledge_docs"], "part": 2},
        metadata={"skill": "question_extractor", "source_line": payload.get("source_line")},
    )


def extract_background_candidates(text: str, job: IngestionJob) -> list[Candidate]:
    candidates: list[Candidate] = []
    for line in text.splitlines():
        match = BACKGROUND_FIELD_RE.match(line.strip())
        if not match:
            continue
        raw_key = match.group("key").strip()
        value = match.group("value").strip()
        if not value:
            continue
        fact_key = normalize_fact_key(raw_key)
        candidates.append(
            Candidate(
                candidate_kind="background",
                title=f"背景信息：{fact_key}",
                summary=value[:120],
                content=value,
                normalized_payload={
                    "topic": topic_for_fact_key(fact_key),
                    "fact_key": fact_key,
                    "fact_value": value,
                    "privacy_level": "normal",
                    "allowed_usage": ["question_personalization", "feedback_personalization"],
                },
                materialization_plan={"tables": ["background_questionnaires", "background_facts"], "requires_user_confirmation": True},
                metadata={"skill": "background_fact_extractor", "raw_key": raw_key, "owner_user_id": job.owner_user_id},
            )
        )
    return candidates[:30]


def extract_knowledge_candidates(text: str, job: IngestionJob, *, avoid_duplicate_question_texts: list[Candidate]) -> list[Candidate]:
    question_texts = {item.content.strip().lower() for item in avoid_duplicate_question_texts if item.candidate_kind == "question"}
    sections = split_sections(text)
    candidates: list[Candidate] = []
    for title, content in sections:
        normalized = content.strip()
        if len(normalized) < 80:
            continue
        if normalized.lower() in question_texts:
            continue
        candidates.append(
            Candidate(
                candidate_kind="knowledge",
                title=title or job.source_file.title,
                summary=first_sentence(normalized),
                content=normalized[:12000],
                normalized_payload={"title": title or job.source_file.title, "content": normalized[:12000]},
                materialization_plan={"tables": ["knowledge_docs", "knowledge_chunks"], "doc_type": "topic_knowledge"},
                metadata={"skill": "knowledge_chunker", "source_file_id": job.source_file.id},
            )
        )
    return candidates[:20]


def split_sections(text: str) -> list[tuple[str | None, str]]:
    lines = text.splitlines()
    sections: list[tuple[str | None, list[str]]] = []
    current_title: str | None = None
    current_lines: list[str] = []
    for line in lines:
        heading = HEADING_RE.match(line)
        if heading:
            if current_lines:
                sections.append((current_title, current_lines))
            current_title = heading.group("title").strip()
            current_lines = []
            continue
        current_lines.append(line)
    if current_lines:
        sections.append((current_title, current_lines))
    if not sections:
        return [(None, text)]
    normalized: list[tuple[str | None, str]] = []
    for title, block_lines in sections:
        block = "\n".join(block_lines).strip()
        if block:
            normalized.append((title, block))
    return normalized or [(None, text)]


def dedupe_candidates(candidates: list[Candidate]) -> list[Candidate]:
    seen: set[tuple[str, str]] = set()
    unique: list[Candidate] = []
    for candidate in candidates:
        key = (candidate.candidate_kind, normalize_text(candidate.content).lower()[:240])
        if key in seen:
            continue
        seen.add(key)
        unique.append(candidate)
    return unique


def normalize_text(text: str) -> str:
    text = strip_nul_bytes(text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def part_from_text(text: str) -> int:
    match = re.search(r"part\s*([123])", text, re.IGNORECASE)
    if match:
        return int(match.group(1))
    return 1


def short_title(text: str, limit: int = 80) -> str:
    cleaned = re.sub(r"\s+", " ", text).strip(" -#")
    if len(cleaned) <= limit:
        return cleaned
    return cleaned[: limit - 1].rstrip() + "..."


def first_sentence(text: str, limit: int = 180) -> str:
    cleaned = re.sub(r"\s+", " ", text).strip()
    if not cleaned:
        return ""
    match = re.search(r"(.+?[。.!?])\s", cleaned)
    sentence = match.group(1) if match else cleaned
    return short_title(sentence, limit)


def normalize_fact_key(raw_key: str) -> str:
    mapping = {
        "name": "name",
        "目标分数": "target_band",
        "target band": "target_band",
        "occupation": "occupation",
        "job": "occupation",
        "职业": "occupation",
        "major": "major",
        "专业": "major",
        "school": "school",
        "学校": "school",
        "hobby": "hobby",
        "兴趣": "hobby",
        "weakness": "weakness",
        "弱项": "weakness",
        "strength": "strength",
        "强项": "strength",
        "city": "city",
        "城市": "city",
    }
    return mapping.get(raw_key.strip().lower(), re.sub(r"\W+", "_", raw_key.strip().lower()).strip("_"))


def topic_for_fact_key(fact_key: str) -> str:
    if fact_key in {"target_band", "weakness", "strength"}:
        return "learning_goal"
    if fact_key in {"occupation", "major", "school"}:
        return "education_work"
    return "personal_profile"


def looks_like_background(text: str) -> bool:
    return any(BACKGROUND_FIELD_RE.match(line.strip()) for line in text.splitlines())


def candidates_to_json(candidates: list[Candidate]) -> str:
    payload = [asdict(candidate) for candidate in candidates]
    return json.dumps(payload, ensure_ascii=False, indent=2, default=str)
